# 채점기 — 고치지 말 것
import pytest

import 시작

openpyxl = pytest.importorskip("openpyxl")
from openpyxl import Workbook


def _시트(행들, 병합=()):
    """행 목록으로 시트를 만들고 병합 범위를 걸어 돌려준다."""
    wb = Workbook()
    ws = wb.active
    for 행 in 행들:
        ws.append(행)
    for r in 병합:
        ws.merge_cells(r)
    return ws


def test_세로_병합이_채워진다():
    ws = _시트(
        [["구분", "1분기", "2분기"],
         ["인프라사업부", 4200000, 3800000],
         [None, 1200000, 1400000],
         ["구매팀", 900000, 1000000]],
        병합=["A2:A3"],
    )
    grid = 시작._fill_merged(ws)
    # 가려진 칸이 왼쪽 위 값으로 채워져야 한다
    assert grid[2][0] == "인프라사업부"
    # 병합과 무관한 칸은 그대로
    assert grid[3][0] == "구매팀"
    assert grid[2][1] == 1200000


def test_가로_병합이_채워진다():
    ws = _시트(
        [["2026년 집계", None, None],
         ["부서", "1분기", "2분기"],
         ["구매팀", 100, 200]],
        병합=["A1:C1"],
    )
    grid = 시작._fill_merged(ws)
    assert grid[0] == ["2026년 집계", "2026년 집계", "2026년 집계"]


def test_빈_시트에서_터지지_않는다():
    wb = Workbook()
    assert 시작._fill_merged(wb.active) == []


def test_병합이_없으면_원본_그대로():
    ws = _시트([["가", 1], ["나", 2]])
    assert 시작._fill_merged(ws) == [["가", 1], ["나", 2]]


def test_A1에서_시작하지_않는_시트도_자리가_맞는다():
    # C3 부터 값을 넣는다. iter_rows 는 1행부터 주므로 grid 는 5행이 되고,
    # 좌표 기준을 잘못 잡으면 엉뚱한 칸을 왼쪽 위로 읽어 병합이 안 채워진다.
    wb = Workbook()
    ws = wb.active
    ws["C3"] = "부서"
    ws["D3"] = "금액"
    ws["C4"] = "구매팀"
    ws["D4"] = 100
    ws["D5"] = 200
    ws.merge_cells("C4:C5")

    grid = 시작._fill_merged(ws)
    assert len(grid) == 5                      # 시트 1~5 행
    assert grid[3][2] == "구매팀"               # C4 (원래 값)
    assert grid[4][2] == "구매팀"               # C5 가 채워졌다


def test_여러_병합이_모두_채워진다():
    ws = _시트(
        [["A", "B", "C"],
         ["가", 1, 2],
         [None, 3, 4],
         ["나", 5, 6],
         [None, 7, 8]],
        병합=["A2:A3", "A4:A5"],
    )
    grid = 시작._fill_merged(ws)
    assert grid[2][0] == "가"
    assert grid[4][0] == "나"


def test_돌려주는_것이_리스트의_리스트다():
    ws = _시트([["가", 1]])
    grid = 시작._fill_merged(ws)
    assert isinstance(grid, list)
    # 튜플이 아니라 리스트여야 한다 — 뒤 단계에서 칸을 바꿔 넣기 때문이다
    assert all(isinstance(row, list) for row in grid)
