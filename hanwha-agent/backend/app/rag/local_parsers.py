from __future__ import annotations

from pathlib import Path

from app.core.logging import get_logger
from app.integrations.ports import ParsedBlock, ParsedDoc

log = get_logger(__name__)


# PDF 를 읽을 때 표를 따로 떼어낼지 정하는 스위치.
# pypdf 는 쪽 단위 글자만 주고 표를 모른다. pdfplumber 는 표를 찾지만 느리다.
# 그래서 기본은 꺼 두고 필요할 때만 .env 에서 켠다.
def pdf_tables_on() -> bool:
    import os

    # 환경변수를 먼저 본다. 켜고 끄기를 실행할 때마다 바꿀 수 있어야 해서다.
    raw = os.environ.get("PDF_TABLES")
    if raw is None:
        from dotenv import dotenv_values

        # 환경변수에 없으면 .env 를 읽는다. 
        # 상대 경로라 프로젝트 루트에서 실행해야 찾는다.
        raw = dotenv_values(".env").get("PDF_TABLES")
    # "true" 만 받지 않는다. 사람이 1, yes, on 중 무엇을 적을지 모른다.
    return str(raw or "").strip().lower() in ("1", "true", "yes", "on")


# PDF 를 쪽 단위로 읽어 ParsedDoc 으로 돌려준다. 표는 못 가져온다.
def parse_pdf(path: Path) -> ParsedDoc:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks: list[ParsedBlock] = []


    for page_no, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if not text:
            # 글자가 0인 쪽은 스캔본(이미지)이다. 예외로 올리지 않고 기록만 남기고 넘어간다 —
            # 16쪽 중 1쪽이 그림이라고 문서 전체를 버릴 이유가 없다.
            log.info("텍스트 없음 (스캔본으로 보임): %s p.%d", path.name, page_no)
            continue
        blocks.append(ParsedBlock("조항", f"p.{page_no}", text))

    return ParsedDoc(blocks=blocks, page_count=len(reader.pages), table_count=0)


# DOCX 를 본문 순서 그대로 읽어 ParsedDoc 으로 돌려준다.
def parse_docx(path: Path) -> ParsedDoc:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    blocks: list[ParsedBlock] = []
    tables = 0

    # doc.paragraphs 와 doc.tables 를 따로 읽으면 문서 순서가 무너진다.
    # body 의 자식을 차례로 훑어야 3번째 문단과 5번째 문단 사이의 표가 그 자리에 남는다.
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]      # 네임스페이스를 떼면 p 또는 tbl 이 남는다
        if tag == "p":
            text = Paragraph(child, doc).text.strip()
            if text:
                blocks.append(ParsedBlock("조항", f"{path.stem}", text))
        elif tag == "tbl":
            table = Table(child, doc)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            tables += 1
            blocks.append(
                ParsedBlock("표", f"표{tables}", _table_to_markdown(rows[0], rows[1:]))
            )

    return ParsedDoc(blocks=blocks, page_count=1, table_count=tables)


# 머리행과 나머지 행을 마크다운 표 한 덩어리로. 형식이 달라도 표의 모양은 하나다.
def _table_to_markdown(headers: list[str], rows: list[list[str]]) -> str:
    head = "| " + " | ".join(headers) + " |"
    sep = "|" + "---|" * len(headers)
    body = ["| " + " | ".join(str(c) for c in r) + " |" for r in rows]
    return "\n".join([head, sep, *body])


# 표의 칸 하나를 마크다운 표에 넣어도 되는 한 줄로 바꾼다.
def _cell(value) -> str:
    if value is None:
        return ""
    # split() 뒤 join 이면 줄바꿈·탭·연속 공백이 한 칸으로 정리된다.
    # 칸 안의 줄바꿈이 남으면 마크다운 표의 행 구분이 깨진다.
    text = " ".join(str(value).split())
    # | 는 마크다운 표의 칸 구분자다. 글자로 쓰려면 이스케이프해야 열이 밀리지 않는다.
    return text.replace("|", "\\|")


# 병합된 칸의 값을 가려진 칸에도 채워 넣는다.
# openpyxl 은 병합 범위의 첫 칸에만 값을 주고 나머지는 None 으로 준다.
# 그대로 두면 '| 인프라사업부 | 4200000 |' 다음 줄이 '|  | 1200000 |' 이 되어
# 1,200,000원이 어느 부서 것인지 사라진다.
def _fill_merged(ws) -> list[list]:
    grid = [list(row) for row in ws.iter_rows(values_only=True)]
    if not grid:
        return grid
    # 병합 범위는 절대 좌표(시트의 몇 행 몇 열)로 오므로 grid 의 0 번째가
    # 시트의 몇 행인지 알아야 자리를 맞출 수 있다.
    #
    # 순서가 중요하다. 위의 iter_rows() 가 1행부터 훑으면서 빈 칸을 만들어 놓기 때문에
    # 그 뒤에 읽는 min_row 는 1 이 된다(C3 부터 쓴 시트도 호출 전에는 3 이었다).
    # grid[0] 이 1행이므로 이때만 자리가 맞는다 — 이 두 줄을 위로 올리면 깨진다.
    r0, c0 = ws.min_row, ws.min_column
    for rng in ws.merged_cells.ranges:
        top_left = grid[rng.min_row - r0][rng.min_col - c0]
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                grid[r - r0][c - c0] = top_left
    return grid


# 슬라이드 안의 표를 마크다운으로. has_text_frame 으로는 안 잡히는 도형이다.
def _pptx_table(shape) -> str:
    rows = [[_cell(cell.text) for cell in row.cells] for row in shape.table.rows]
    # 칸이 전부 빈 표는 버린다. 레이아웃용으로 넣은 빈 표가 검색 결과에 끼지 않게.
    if not any(any(row) for row in rows):
        return ""
    return _table_to_markdown(rows[0], rows[1:])


# 계산값이 저장돼 있지 않은 수식 칸을 표시한다.
# data_only=True 로 읽으면 수식 칸은 엑셀이 저장할 때 계산해 둔 값으로 온다.
# 그런데 파이썬으로 만든 파일처럼 그 값이 없으면 None 이 와서 빈칸이 된다.
# 빈칸으로 두면 '값이 없는 것'과 '계산이 안 된 것'을 구별할 수 없다.
def _fill_formulas(grid: list[list], ws, ws_formula) -> list[list]:
    r0, c0 = ws.min_row, ws.min_column
    for i, row in enumerate(grid):
        for j, value in enumerate(row):
            if value is not None:
                continue
            # 같은 시트를 data_only 없이 한 번 더 읽어 둔 것이 ws_formula 다.
            # 거기서는 수식 문자열 자체("=B3*2")가 온다.
            formula = ws_formula.cell(row=r0 + i, column=c0 + j).value
            if isinstance(formula, str) and formula.startswith("="):
                row[j] = f"{formula} (계산값 없음)"
    return grid


# 한 시트에 표가 여러 개 있을 때 빈 줄을 경계로 쪼갠다.
# 사람은 엑셀 한 시트에 표를 여러 개 두고 빈 줄로 구분한다.
# 통째로 한 표로 만들면 머리행이 하나뿐인 뒤죽박죽 표가 된다.
def _split_regions(grid: list[list], first_row: int) -> list[tuple[int, list[list]]]:

    regions: list[tuple[int, list[list]]] = []
    for i, row in enumerate(grid):
        if all(c is None for c in row):
            continue                                   # 빈 줄은 경계이므로 담지 않는다
        # 바로 앞 줄에 이어지는 행이면 같은 표, 건너뛰었으면 새 표다.
        if regions and regions[-1][0] + len(regions[-1][1]) == first_row + i:
            regions[-1][1].append(row)
        else:
            regions.append((first_row + i, [row]))
    # 표가 둘 이상이면 각 표의 폭을 그 표가 실제로 쓰는 만큼으로 줄인다.
    # 시트 전체 폭을 쓰면 좁은 표에 빈 열이 줄줄이 붙는다.
    if len(regions) > 1:
        for _, rows in regions:
            width = max(max((j + 1 for j, c in enumerate(r) if c is not None), default=0) for r in rows)
            rows[:] = [r[:width] for r in rows]
    return regions


# 표 하나를 마크다운으로. 제목행과 두 줄 머리행을 가려낸다.
def _xlsx_table(ws, top: int, rows: list[list[str]]) -> str:
    # 그 행에서 가로로 병합이 걸려 있는지 본다. 제목행·묶음 머리행의 표시다.
    def merged_across(row_no: int) -> bool:
        return any(r.min_row == row_no and r.max_col > r.min_col for r in ws.merged_cells.ranges)

    title = ""
    # 첫 행이 가로로 병합돼 있고 모든 칸이 같은 값이면 머리행이 아니라 표 제목이다.
    # (_fill_merged 가 병합된 칸을 같은 값으로 채워 둔 덕에 set() 크기로 알 수 있다)
    if len(rows) > 2 and len(set(rows[0])) == 1 and len(rows[0]) > 1 and merged_across(top):
        title, rows, top = rows[0][0], rows[1:], top + 1
    header, body = rows[0], rows[1:]
    # 머리행이 두 줄로 묶인 표("1분기 / 매출·비용")는 두 줄을 합쳐 한 줄 머리행으로 만든다.
    if body and merged_across(top):
        second = body[0]
        header = [a if a == b or not b else (b if not a else f"{a} {b}") for a, b in zip(header, second)]
        body = body[1:]
    table = _table_to_markdown(header, body)
    return f"{title}\n\n{table}" if title else table


# 그룹으로 묶인 도형 안쪽까지 들어가 도형을 하나씩 내놓는다.
# slide.shapes 만 돌면 그룹은 하나로 보이고 그 안의 글자·표는 통째로 빠진다.
def _walk_shapes(shapes):
    from pptx.shapes.group import GroupShape

    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _walk_shapes(shape.shapes)      # 그룹이면 한 겹 더 들어간다
        else:
            yield shape


# 슬라이드의 차트를 표로 바꾼다. 차트도 숫자가 든 자료라 검색 대상이다.
def _pptx_chart(shape) -> str:
    chart = shape.chart
    try:
        categories = [_cell(c) for c in chart.plots[0].categories]
    except (IndexError, AttributeError, TypeError):
        # 차트 종류에 따라 categories 가 없다. 못 읽는 차트는 버리고 넘어간다 —
        # 차트 하나 때문에 슬라이드 전체 파싱이 멈추면 안 된다.
        return ""
    series = list(chart.series)
    if not categories or not series:
        return ""

    # 엑셀 숫자는 전부 float 로 온다. 4200000.0 을 4200000 으로 적어야 읽힌다.
    def num(v) -> str:
        return _cell(int(v) if isinstance(v, float) and v.is_integer() else v)

    header = ["항목", *(_cell(s.name) for s in series)]
    body = [[cat, *(num(s.values[i]) if i < len(s.values) else "" for s in series)]
            for i, cat in enumerate(categories)]
    table = _table_to_markdown(header, body)
    title = chart.chart_title.text_frame.text.strip() if chart.has_title else ""
    return f"{title}\n\n{table}" if title else table


# PDF 를 pdfplumber 로 읽어 표는 표 블록, 나머지는 조항 블록으로 만든다.
# pypdf 로는 표가 0개였던 문서에서 7개가 나온다.
def parse_pdf_tables(path: Path) -> ParsedDoc:
    import pdfplumber

    blocks: list[ParsedBlock] = []
    n_tables = 0
    with pdfplumber.open(str(path)) as pdf:
        page_count = len(pdf.pages)
        for page_no, page in enumerate(pdf.pages, 1):
            found = page.find_tables()
            boxes = [t.bbox for t in found]       # (x0, top, x1, bottom)

            # 글자 하나가 어느 표 상자에도 안 들어가면 참. 본문만 남기는 체다.
            # 이게 없으면 표 안의 글자가 본문에도 한 번 더 실려 같은 내용이 두 벌이 된다.
            # boxes=boxes 로 기본값에 묶어 두는 이유는 반복문 변수를 늦게 읽지 않게 하려는 것이다.
            def outside(obj, boxes=boxes) -> bool:
                x0, top = obj.get("x0", 0), obj.get("top", 0)
                x1, bottom = obj.get("x1", 0), obj.get("bottom", 0)
                return not any(b[0] <= x0 and x1 <= b[2] and b[1] <= top and bottom <= b[3]
                               for b in boxes)

            # 표가 없는 쪽은 걸러낼 것이 없으니 page 를 그대로 쓴다.
            text = ((page.filter(outside) if boxes else page).extract_text() or "").strip()
            if text:
                blocks.append(ParsedBlock("조항", f"p.{page_no}", text))
            for k, table in enumerate(found, 1):
                rows = [[_cell(c) for c in row] for row in table.extract()]
                if not any(any(row) for row in rows):
                    continue                     # 칸이 전부 빈 표는 버린다
                blocks.append(ParsedBlock("표", f"p.{page_no} · 표{k}",
                                          _table_to_markdown(rows[0], rows[1:])))
                n_tables += 1
            # 글자도 표도 없으면 그 쪽은 그림이다.
            if not text and not found:
                log.info("텍스트 없음 (스캔본으로 보임): %s p.%d", path.name, page_no)

    return ParsedDoc(blocks=blocks, page_count=page_count, table_count=n_tables)


# HWPX 는 ZIP 안의 XML 이다. 한글 전용 라이브러리 없이 표준 모듈로 읽는다.
def parse_hwpx(path: Path) -> ParsedDoc:
    import re
    import zipfile
    from xml.etree import ElementTree as ET

    # 한컴이 정한 네임스페이스. 태그 이름 앞에 이게 붙어 있어야 찾는다.
    ns_p = "{http://www.hancom.co.kr/hwpml/2011/paragraph}"
    blocks: list[ParsedBlock] = []
    tables = 0

    with zipfile.ZipFile(path) as zf:
        # 본문은 section0.xml, section1.xml … 로 나뉘어 있다. 순서대로 읽어야 문서 순서가 산다.
        sections = sorted(n for n in zf.namelist()
                          if re.match(r"Contents/section\d+\.xml$", n))
        for name in sections:
            root = ET.fromstring(zf.read(name))

            # findall 은 바로 아래 자식만 본다 — 문단을 세는 자리라 그게 맞다.
            for para in root.findall(f"{ns_p}p"):
                table_el = para.find(f".//{ns_p}tbl")
                if table_el is not None:
                    tables += 1
                    rows: list[list[str]] = []
                    for tr in table_el.findall(f"{ns_p}tr"):
                        cells = []
                        for tc in tr.findall(f"{ns_p}tc"):
                            # iter 는 깊이 상관없이 훑는다 — 칸 안의 글자를 다 모으는 자리다.
                            cells.append("".join(
                                t.text or "" for t in tc.iter(f"{ns_p}t")
                            ).strip())
                        rows.append(cells)
                    if rows:
                        blocks.append(ParsedBlock(
                            "표", f"표{tables}", _table_to_markdown(rows[0], rows[1:])
                        ))
                    continue                       # 표였으면 문단 처리는 건너뛴다
                text = "".join(t.text or "" for t in para.iter(f"{ns_p}t")).strip()
                if text:
                    blocks.append(ParsedBlock("조항", path.stem, text))

    return ParsedDoc(blocks=blocks, page_count=1, table_count=tables)


# XLSX 를 읽는다. 시트 하나에 표가 여러 개일 수 있다.
def parse_xlsx(path: Path) -> ParsedDoc:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True)
    # 같은 파일을 data_only 없이 한 번 더 연다. 수식 문자열을 보려면 이쪽이 필요하다.
    wb_formula = load_workbook(str(path))
    blocks: list[ParsedBlock] = []

    for ws in wb.worksheets:
        # 순서가 중요하다. 병합을 먼저 채워야 _split_regions 가 빈 줄을 제대로 알아보고,
        # _xlsx_table 이 제목행을 set() 크기로 알아볼 수 있다.
        grid = _fill_formulas(_fill_merged(ws), ws, wb_formula[ws.title])
        regions = _split_regions(grid, ws.min_row)
        for k, (top, region) in enumerate(regions, 1):
            rows = [["" if c is None else str(c) for c in row] for row in region]
            # 표가 하나면 시트 이름만, 여럿이면 번호를 붙여 어느 표인지 남긴다.
            locator = ws.title if len(regions) == 1 else f"{ws.title} · 표{k}"
            blocks.append(ParsedBlock("표", locator, _xlsx_table(ws, top, rows)))  # ◀ 추가 끝

    return ParsedDoc(
        blocks=blocks, page_count=len(wb.worksheets), table_count=len(blocks)
    )


# PPTX 를 읽는다. 슬라이드 하나가 블록 하나이고 표·차트는 따로 떼어낸다.
def parse_pptx(path: Path) -> ParsedDoc:
    from pptx import Presentation

    prs = Presentation(str(path))
    blocks: list[ParsedBlock] = []
    n_tables = 0

    for page_no, slide in enumerate(prs.slides, 1):
        # has_text_frame 으로 거른다. 이미지·선·차트에는 글자 틀이 없어서
        # shape.text 를 바로 읽으면 거기서 예외가 난다.
        lines = [shape.text.strip() for shape in _walk_shapes(slide.shapes)
                 if shape.has_text_frame and shape.text.strip()]
        # 발표자 노트는 slide.shapes 에 없다. 화면에 안 보이는 글이라 따로 가져온다.
        if slide.has_notes_slide:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                lines.append(f"[발표자 노트] {note}")
        if lines:
            blocks.append(ParsedBlock("조항", f"슬라이드 {page_no}", "\n".join(lines)))
        shapes = list(_walk_shapes(slide.shapes))
        # 빈 문자열을 돌려준 표·차트는 걸러낸다(읽을 것이 없던 것들).
        tables = [md for md in (_pptx_table(s) for s in shapes if s.has_table) if md]
        charts = [md for md in (_pptx_chart(s) for s in shapes if s.has_chart) if md]
        for k, md in enumerate(tables, 1):
            blocks.append(ParsedBlock("표", f"슬라이드 {page_no} · 표{k}", md))
        for k, md in enumerate(charts, 1):
            blocks.append(ParsedBlock("표", f"슬라이드 {page_no} · 차트{k}", md))
        n_tables += len(tables) + len(charts)

    return ParsedDoc(blocks=blocks, page_count=len(prs.slides), table_count=n_tables)


# 못 읽는 형식에 "그럼 어떻게 하라"를 적어 둔다. 막을 때 이유만 말하지 않는다.
_CONVERT_HINT = {
    ".hwp": "한글에서 열고 [다른 이름으로 저장] → HWPX 또는 PDF 로 내보내세요. "
            "구 .hwp 는 한컴 독자 바이너리라 열지 않고는 읽을 방법이 없습니다",
    ".doc": "Word 에서 열고 .docx 로 저장하세요",
    ".ppt": "PowerPoint 에서 열고 .pptx 로 저장하세요",
    ".xls": "Excel 에서 열고 .xlsx 로 저장하세요",
    ".hwt": "한글 서식 파일입니다. 내용이 든 .hwpx 문서를 넣으세요",
    ".zip": "압축을 풀고 안의 문서를 한 건씩 넣으세요",
}

# 확장자 → 파서. 여기 없으면 읽지 않는다. 형식을 늘리는 일은 이 표에 한 줄 더하는 일이다.
HANDLERS = {
    ".docx": parse_docx,
    ".pdf": parse_pdf,
    ".hwpx": parse_hwpx,
    ".xlsx": parse_xlsx,
    ".pptx": parse_pptx,
}

# 확장자를 보고 알맞은 파서로 넘겨주는 진입
# 파일 하나를 읽어 ParsedDoc 으로. 끝나는 길이 셋이다 
# — 정상 / None / ValidationFailed.
def parse_local(path: str | Path) -> ParsedDoc | None:
    from app.core.exceptions import ValidationFailed

    p = Path(path)

    ext = p.suffix.lower()

    handler = HANDLERS.get(ext)
    # PDF 는 스위치에 따라 파서가 바뀐다. 분기표는 그대로 두고 여기서 갈아 끼운다.
    if handler is parse_pdf and pdf_tables_on():
        handler = parse_pdf_tables
    if handler is None:
        if ext in _CONVERT_HINT:
            raise ValidationFailed(
                f"읽을 수 없는 형식입니다: {ext}",
                detail=_CONVERT_HINT[ext],
            )
        raise ValidationFailed(
            f"지원하지 않는 형식입니다: {ext}",
            detail="핸들러를 추가하면 지원할 수 있습니다",
        )

    try:
        doc = handler(p)
    except Exception as exc:
        # 파서가 터진 것은 로그만 남기고 None 을 돌려준다. 올린 사람이 고칠 수 있는 일이 아니다.
        log.warning("로컬 파싱 실패 (%s): %s", p.name, exc)
        return None

    # 블록이 0개면 성공으로 넘기지 않는다. 글자 없는 문서를 적재하면 검색에서 영영 안 나온다.
    if not doc.blocks:
        raise ValidationFailed(
            f"문서에서 글자를 찾지 못했습니다: {p.name}",
            detail="스캔본(이미지)일 수 있습니다. .env 의 UPSTAGE_PARSE_OCR 를 force 로 "
                   "두고 실동작 모드로 다시 적재해 보세요",
        )

    return doc
